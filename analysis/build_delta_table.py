# -*- coding: utf-8 -*-
"""
build_delta_table.py
--------------------------------
미팅 자료용 "지표 x rate구간 델타" 표를 엑셀로 자동 생성. 그래프를 눈대중으로
읽어서 표를 만들면 숫자가 부정확해질 위험이 있어서, 실제
results/<ds>/all_results_test_only.csv에서 정확한 값을 가져와 계산한다.

표 구조 (사용자 요청 그대로):
  - 세로: 모델 5개(기본 USAD/AutoEncoder/LSTMAD/TimesNet/Donut) + 맨 아래 평균행
  - 가로: 지표 4개(AUC-ROC/AUC-PR/VUS-ROC/VUS-PR) x 각 지표마다
    [0%(baseline), 0→1%, 0→5%, 0→10%, 0→20%] 5개 서브컬럼
  - baseline은 rate=0%일 때의 실측값 그대로, 나머지 4개는
    (해당 rate 값 - rate=0% 값)의 델타
  - 데이터셋(예: smd, psm)마다 별도 시트로 생성, 여러 개 한 번에 가능

원본 rate별 실측값은 "raw_<dataset>" 시트에 따로 넣어두고, 보이는 표의 델타
칸은 전부 그 시트를 참조하는 수식(예: =raw_smd!C3-raw_smd!B3)으로 계산 —
숫자를 파이썬에서 미리 계산해서 박아넣지 않음. 평균행도 AVERAGE 수식.

사용법:
    python build_delta_table.py --datasets smd psm --out delta_table.xlsx
    python build_delta_table.py --datasets smd psm swat --models USAD Donut AutoEncoder --out delta_table.xlsx
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config

METRICS = ["AUC-ROC", "AUC-PR", "VUS-ROC", "VUS-PR"]
RATE_COLS = [0, 1, 5, 10, 20]  # 0은 baseline, 나머지는 "0%->해당rate%" 델타
FONT_NAME = "Arial"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
SUBHEADER_FONT = Font(name=FONT_NAME, bold=True)
AVG_FILL = PatternFill("solid", fgColor="FFF2CC")
AVG_FONT = Font(name=FONT_NAME, bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RAW_FONT = Font(name=FONT_NAME, color="0000FF")  # 하드코딩 입력값 = 파란색 관례


def write_raw_sheet(wb: Workbook, dataset_key: str, models: list, df: pd.DataFrame):
    """rate별 실측 metric 값을 그대로 넣는 시트 (표의 델타 수식이 참조하는 원본)."""
    ws = wb.create_sheet(f"raw_{dataset_key}")
    ws["A1"] = f"{dataset_key} — rate별 실측값 (all_results_test_only.csv 평균, model별)"
    ws["A1"].font = Font(name=FONT_NAME, bold=True)

    row = 3
    cell_map = {}  # (model, metric) -> {rate: cell_ref}
    for metric in METRICS:
        ws.cell(row=row, column=1, value=metric).font = Font(name=FONT_NAME, bold=True)
        header_row = row + 1
        ws.cell(row=header_row, column=1, value="model")
        for j, rate in enumerate(RATE_COLS):
            ws.cell(row=header_row, column=2 + j, value=f"rate_{rate}")
        for i, model_name in enumerate(models):
            r = header_row + 1 + i
            ws.cell(row=r, column=1, value=model_name)
            grp = df[(df["model"] == model_name)]
            for j, rate in enumerate(RATE_COLS):
                val = grp[grp["contamination_rate"] == rate / 100][metric].mean()
                cell = ws.cell(row=r, column=2 + j,
                                value=None if pd.isna(val) else round(float(val), 4))
                cell.font = RAW_FONT
                cell_map[(model_name, metric, rate)] = cell.coordinate
        row = header_row + len(models) + 2

    return cell_map


def write_delta_sheet(wb: Workbook, dataset_key: str, models: list, cell_map: dict):
    ws = wb.create_sheet(f"delta_{dataset_key}", 0)
    raw_sheet = f"raw_{dataset_key}"

    # 1행: 데이터셋 라벨 / 2행: metric 그룹 병합 헤더 / 3행: rate 서브헤더
    ws["A1"] = f"{dataset_key} — Contamination Rate에 따른 지표 변화 (baseline=rate 0%)"
    ws["A1"].font = Font(name=FONT_NAME, bold=True, size=12)
    ws.merge_cells("A1:U1")

    ws.cell(row=2, column=1, value="model")
    ws.cell(row=3, column=1, value="")
    ws.cell(row=2, column=1).font = HEADER_FONT
    ws.cell(row=2, column=1).fill = HEADER_FILL
    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)

    sub_labels = ["0%", "0→1%", "0→5%", "0→10%", "0→20%"]
    col = 2
    metric_start_col = {}
    for metric in METRICS:
        metric_start_col[metric] = col
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + len(sub_labels) - 1)
        c = ws.cell(row=2, column=col, value=metric)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
        for k, label in enumerate(sub_labels):
            sc = ws.cell(row=3, column=col + k, value=label)
            sc.font = SUBHEADER_FONT
            sc.fill = SUBHEADER_FILL
            sc.alignment = Alignment(horizontal="center")
        col += len(sub_labels)

    # 모델별 행
    first_data_row = 4
    for i, model_name in enumerate(models):
        r = first_data_row + i
        ws.cell(row=r, column=1, value=model_name).font = Font(name=FONT_NAME, bold=True)
        for metric in METRICS:
            base_col = metric_start_col[metric]
            base_ref = f"'{raw_sheet}'!{cell_map[(model_name, metric, 0)]}"
            # 0% baseline 칸: 원본값 그대로 참조
            ws.cell(row=r, column=base_col, value=f"={base_ref}").number_format = "0.0%"
            # 나머지 4칸: (해당 rate 값 - baseline) 델타 수식
            for k, rate in enumerate(RATE_COLS[1:], start=1):
                rate_ref = f"'{raw_sheet}'!{cell_map[(model_name, metric, rate)]}"
                cell = ws.cell(row=r, column=base_col + k, value=f"={rate_ref}-{base_ref}")
                cell.number_format = "+0.0%;-0.0%"

    # 평균행
    avg_row = first_data_row + len(models)
    ws.cell(row=avg_row, column=1, value="평균").font = AVG_FONT
    ws.cell(row=avg_row, column=1).fill = AVG_FILL
    last_data_row = avg_row - 1
    for metric in METRICS:
        base_col = metric_start_col[metric]
        for k in range(len(sub_labels)):
            col_letter = get_column_letter(base_col + k)
            cell = ws.cell(row=avg_row, column=base_col + k,
                            value=f"=AVERAGE({col_letter}{first_data_row}:{col_letter}{last_data_row})")
            cell.number_format = "0.0%" if k == 0 else "+0.0%;-0.0%"
            cell.font = AVG_FONT
            cell.fill = AVG_FILL

    # 테두리 + 열너비
    max_col = 1 + len(METRICS) * len(sub_labels)
    for r in range(2, avg_row + 1):
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).border = BORDER
    ws.column_dimensions["A"].width = 14
    for c in range(2, max_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 9

    ws.freeze_panes = "B4"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--datasets", nargs="+", required=True)
    ap.add_argument("--models", nargs="+",
                     default=["USAD", "AutoEncoder", "LSTMAD", "TimesNet", "Donut"])
    ap.add_argument("--out", default="delta_table.xlsx")
    args = ap.parse_args()

    wb = Workbook()
    wb.remove(wb.active)  # 기본 빈 시트 제거

    any_written = False
    for ds in args.datasets:
        csv_path = Path(config.RESULTS_DIR) / ds / "all_results_test_only.csv"
        if not csv_path.exists():
            print(f"⚠ {csv_path} 없음, {ds} 스킵 (analyze_tsbad_results.py --dataset {ds} 먼저 실행)")
            continue
        df = pd.read_csv(csv_path)
        missing_models = [m for m in args.models if m not in df["model"].unique()]
        if missing_models:
            print(f"⚠ [{ds}] csv에 없는 모델: {missing_models} (실제: {sorted(df['model'].unique())})")
        models = [m for m in args.models if m in df["model"].unique()]
        if not models:
            print(f"⚠ [{ds}] 지정한 모델이 하나도 없어 스킵")
            continue
        cell_map = write_raw_sheet(wb, ds, models, df)
        write_delta_sheet(wb, ds, models, cell_map)
        any_written = True
        print(f"[{ds}] {len(models)}개 모델 x {len(METRICS)}개 지표 표 작성 완료")

    if not any_written:
        print("작성된 시트가 없어 저장하지 않음.")
        return

    wb.save(args.out)
    print(f"\n저장: {args.out}")
    print("※ 수식만 써있고 캐시된 값은 없는 상태 — 반드시 recalc 필요:")
    print(f"   python recalc.py {args.out}   (xlsx 스킬의 scripts/recalc.py)")


if __name__ == "__main__":
    main()